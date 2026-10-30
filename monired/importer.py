"""Import the network inventory spreadsheet: one sheet becomes one zone."""

from __future__ import annotations

import ipaddress
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .db import Inventory

# Normalized header text -> inventory field.
HEADERS = {
    "ip": "ip",
    "codigo del salon": "code",
    "codigo": "code",
    "sede": "sede",
    "ciclo": "ciclo",
    "cargo": "name",
    "nombre": "name",
    "estado": "estado",
    "observaciones": "notes",
}


@dataclass
class DeviceRow:
    grp: str
    code: str
    name: str
    ip: str | None
    monitored: bool
    notes: str


@dataclass
class ZoneData:
    name: str
    devices: list[DeviceRow] = field(default_factory=list)


@dataclass
class ImportReport:
    zones: dict[str, int] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    @property
    def total(self) -> int:
        return sum(self.zones.values())


def fold_text(value) -> str:
    """Lowercase, accent-free, single-spaced text for tolerant comparisons."""
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return " ".join(text.lower().split())


def _text(value) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _is_ipv4(value) -> bool:
    try:
        ipaddress.IPv4Address(_text(value))
        return True
    except ValueError:
        return False


def read_workbook(path: str | Path) -> tuple[list[ZoneData], list[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    zones, warnings = [], []
    for ws in wb.worksheets:
        rows = [r for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
        zone = ZoneData(ws.title.strip())
        columns = _find_header(rows)
        if columns is None:
            _read_headerless(zone, rows, warnings)
        else:
            header_index, cols = columns
            _read_with_header(zone, rows[header_index + 1 :], cols, warnings)
        if zone.devices:
            zones.append(zone)
    wb.close()
    return zones, warnings


def _find_header(rows) -> tuple[int, dict[str, int]] | None:
    for i, row in enumerate(rows[:10]):
        cols = {}
        for j, cell in enumerate(row):
            key = HEADERS.get(fold_text(cell))
            if key and key not in cols:
                cols[key] = j
        if "ip" in cols:
            return i, cols
    return None


def _read_with_header(zone: ZoneData, rows, cols: dict[str, int], warnings: list[str]) -> None:
    def get(row, key) -> str:
        j = cols.get(key)
        return _text(row[j]) if j is not None and j < len(row) else ""

    for row in rows:
        code, name, raw_ip = get(row, "code"), get(row, "name"), get(row, "ip")
        if not (code or name or raw_ip):
            continue
        sede, ciclo = get(row, "sede"), get(row, "ciclo")
        grp = " / ".join(p for p in (sede, ciclo) if p)
        estado = fold_text(get(row, "estado"))
        zone.devices.append(
            DeviceRow(
                grp=grp,
                code=code,
                name=name,
                ip=_ip_or_warn(raw_ip, zone.name, code or name, warnings),
                monitored=estado != "no operativo",
                notes=get(row, "notes"),
            )
        )


def _read_headerless(zone: ZoneData, rows, warnings: list[str]) -> None:
    """Sheets without headers (e.g. static devices): a text label next to an IP."""
    for row in rows:
        ip_cells = [c for c in row if _is_ipv4(c)]
        if not ip_cells:
            continue
        label = next((_text(c) for c in row if _text(c) and not _is_ipv4(c)), "")
        zone.devices.append(
            DeviceRow(grp="", code="", name=label, ip=_ip_or_warn(ip_cells[0], zone.name, label, warnings),
                      monitored=True, notes="")
        )


def _ip_or_warn(raw, zone: str, who: str, warnings: list[str]) -> str | None:
    text = _text(raw)
    if _is_ipv4(text):
        return str(ipaddress.IPv4Address(text))
    warnings.append(f"{zone} / {who}: IP {text or '(vacía)'!s} no válida, se importa sin IP.")
    return None


def import_workbook(inv: Inventory, path: str | Path, replace: bool = False) -> ImportReport:
    if not inv.is_empty() and not replace:
        raise ValueError("La base ya tiene datos. Usa --reemplazar para borrarla e importar de nuevo.")
    zones, warnings = read_workbook(path)
    report = ImportReport(warnings=warnings)
    if replace:
        inv.clear()
    for zone in zones:
        zone_id = inv.add_zone(zone.name)
        for d in zone.devices:
            inv.add_device(zone_id, grp=d.grp, code=d.code, name=d.name, ip=d.ip,
                           monitored=d.monitored, notes=d.notes)
        report.zones[zone.name] = len(zone.devices)
    for ip, devices in inv.duplicate_ips().items():
        who = ", ".join(f"{d.zone_name} / {d.label}" for d in devices)
        report.warnings.append(f"IP repetida {ip}: {who}")
    return report
