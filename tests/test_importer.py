import pytest
from openpyxl import Workbook

from monired.cli import main
from monired.importer import import_workbook


@pytest.fixture
def workbook(tmp_path):
    """A small copy of the real spreadsheet's layouts."""
    wb = Workbook()
    salones = wb.active
    salones.title = "Salones"
    salones.append([])
    salones.append([None, "Ciclo", "Código del salón", "IP", "Estado", "Observaciones", "#", None,
                    "Total de Equipos Masters"])
    salones.append([None, "Roja", "R1-1", "192.168.10.54", "No Operativo", "Oficina de ejemplo", 0, None, 12])
    salones.append([None, "Roja", "R1-2", "192.168.10.55", "Operativo", None, 1])
    salones.append([])

    coord = wb.create_sheet("Coordinaciones")
    coord.append(["Sede", "Ciclo", "Cargo", "Código", "IP"])
    coord.append(["Roja", "Masters", "Psicologa ", "R-PSICO", "192.168.5.77"])

    oficinas = wb.create_sheet("Oficinas")
    oficinas.append(["Sede", "Cargo", "Código", "IP"])
    oficinas.append(["Principal", "Auxiliar Soporte", "P-ASISTENTETI", "-"])

    estaticas = wb.create_sheet("Estáticas")
    estaticas.append([])
    estaticas.append([None, "SERVIDOR SONIDO", "192.168.10.55"])

    wb.create_sheet("Vacía")
    path = tmp_path / "red.xlsx"
    wb.save(path)
    return path


def test_import_maps_every_sheet_layout(inv, workbook):
    report = import_workbook(inv, workbook)

    assert report.zones == {"Salones": 2, "Coordinaciones": 1, "Oficinas": 1, "Estáticas": 1}
    rows = {(d.zone_name, d.code or d.name): d for d in inv.devices()}

    r11 = rows[("Salones", "R1-1")]
    assert (r11.grp, r11.ip, r11.monitored, r11.notes) == ("Roja", "192.168.10.54", False, "Oficina de ejemplo")
    assert rows[("Salones", "R1-2")].monitored is True

    psico = rows[("Coordinaciones", "R-PSICO")]
    assert (psico.grp, psico.name, psico.ip, psico.monitored) == ("Roja / Masters", "Psicologa", "192.168.5.77", True)

    assert rows[("Oficinas", "P-ASISTENTETI")].ip is None
    assert rows[("Estáticas", "SERVIDOR SONIDO")].ip == "192.168.10.55"

    assert report.warnings == [
        "Oficinas / P-ASISTENTETI: IP - no válida, se importa sin IP.",
        "IP repetida 192.168.10.55: Salones / R1-2, Estáticas / SERVIDOR SONIDO",
    ]


def test_import_refuses_a_non_empty_database_unless_replacing(inv, workbook):
    import_workbook(inv, workbook)
    inv.add_zone("Manual")
    with pytest.raises(ValueError, match="--reemplazar"):
        import_workbook(inv, workbook)

    import_workbook(inv, workbook, replace=True)
    assert [z.name for z in inv.zones()] == ["Salones", "Coordinaciones", "Oficinas", "Estáticas"]
    assert len(inv.devices()) == 5


def test_cli_import(tmp_path, workbook, capsys):
    db = tmp_path / "red.db"
    assert main(["--db", str(db), "importar", str(workbook)]) == 0
    out = capsys.readouterr().out
    assert "Importados 5 equipos en 4 zonas" in out
    assert "IP repetida 192.168.10.55" in out

    assert main(["--db", str(db), "importar", str(workbook)]) == 1
    assert "--reemplazar" in capsys.readouterr().err

    assert main(["--db", str(db), "importar", str(tmp_path / "nope.xlsx")]) == 1
    assert "No existe el archivo" in capsys.readouterr().err


def test_replace_with_clashing_sheet_names_keeps_the_database(inv, workbook, tmp_path):
    import_workbook(inv, workbook)
    before = [(d.zone_name, d.code, d.name, d.ip) for d in inv.devices()]

    wb = Workbook()
    wb.active.title = "Zona A"
    wb.active.append(["Código", "IP"])
    wb.active.append(["A-1", "10.0.0.1"])
    other = wb.create_sheet("ZONA A ")
    other.append(["Código", "IP"])
    other.append(["A-2", "10.0.0.2"])
    clashing = tmp_path / "clash.xlsx"
    wb.save(clashing)

    with pytest.raises(ValueError, match="'Zona A' y 'ZONA A'"):
        import_workbook(inv, clashing, replace=True)
    assert [(d.zone_name, d.code, d.name, d.ip) for d in inv.devices()] == before


def test_cli_reports_clashing_sheet_names(tmp_path, capsys):
    wb = Workbook()
    wb.active.title = "Zona A"
    wb.active.append(["Código", "IP"])
    wb.active.append(["A-1", "10.0.0.1"])
    other = wb.create_sheet("zona a ")
    other.append(["Código", "IP"])
    other.append(["A-2", "10.0.0.2"])
    path = tmp_path / "clash.xlsx"
    wb.save(path)

    assert main(["--db", str(tmp_path / "red.db"), "importar", str(path)]) == 1
    err = capsys.readouterr().err
    assert "Las hojas 'Zona A' y 'zona a' son la misma zona" in err
    assert "--reemplazar" not in err
